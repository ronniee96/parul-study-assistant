"""
Document Processing Service
Handles text extraction from supported file formats (.pdf, .pptx, .docx, .txt, and images when OCR is installed).
"""

import io
import os
import shutil
import hashlib
import logging
from typing import Dict, Any

logger = logging.getLogger(__name__)

SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt", ".jpg", ".jpeg", ".png", ".bmp", ".tiff"}
UNSUPPORTED_LEGACY_BINARY_EXTENSIONS = {".doc", ".ppt"}


class DocumentProcessor:
    """Service for extracting text from supported document formats with provenance metadata."""

    @staticmethod
    def get_ocr_capabilities() -> Dict[str, Any]:
        """Inspect host runtime for local OCR binaries (tesseract, pdftoppm)."""
        tesseract_path = shutil.which("tesseract")
        pdftoppm_path = shutil.which("pdftoppm")
        return {
            "tesseract_installed": bool(tesseract_path),
            "pdftoppm_installed": bool(pdftoppm_path),
            "image_ocr_available": bool(tesseract_path),
            "scanned_pdf_ocr_available": bool(tesseract_path and pdftoppm_path),
        }

    @staticmethod
    def extract_text_from_pdf(file_content: bytes) -> str:
        """Extract text from a digital PDF file using pdfplumber and PyPDF2."""
        if b"%PDF-" not in file_content[:1024]:
            raise ValueError("Invalid PDF file: missing %PDF- header")

        # Try pdfplumber first
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                text = ""
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                if text.strip():
                    return text.strip()
        except Exception as e:
            logger.warning(f"pdfplumber extraction failed, falling back to PyPDF2: {e}")

        # Fallback to PyPDF2
        try:
            import PyPDF2
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            text = ""
            for page in pdf_reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
            return text.strip()
        except Exception as e2:
            logger.error(f"Both PDF extraction methods failed: {e2}")
            raise ValueError(f"Unable to extract text from PDF: {e2}")

    @staticmethod
    def extract_text_from_pptx(file_content: bytes) -> str:
        """Extract text from an OpenXML PowerPoint (.pptx) file."""
        try:
            import pptx
            presentation = pptx.Presentation(io.BytesIO(file_content))
            text = ""

            for slide_num, slide in enumerate(presentation.slides):
                text += f"\n--- Slide {slide_num + 1} ---\n"
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text += shape.text + "\n"

            return text.strip()
        except Exception as e:
            logger.error(f"Failed to extract text from PPTX: {e}")
            raise ValueError(f"Unable to extract text from PowerPoint (.pptx): {e}")

    @staticmethod
    def extract_text_from_docx(file_content: bytes) -> str:
        """Extract text from an OpenXML Word (.docx) document."""
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_content))
            text = ""
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text += paragraph.text + "\n"
            return text.strip()
        except Exception as e:
            logger.error(f"Failed to extract text from DOCX: {e}")
            raise ValueError(f"Unable to extract text from Word (.docx) document: {e}")

    @staticmethod
    def extract_text_from_image(file_content: bytes) -> str:
        """Extract text from an image using Tesseract OCR when the binary is installed."""
        if not shutil.which("tesseract"):
            raise RuntimeError(
                "Image OCR is unavailable: the 'tesseract' binary is not installed in this environment."
            )
        try:
            from PIL import Image
            import pytesseract
            image = Image.open(io.BytesIO(file_content))
            text = pytesseract.image_to_string(image, timeout=3)
            return text.strip()
        except Exception as e:
            logger.error(f"Failed to extract text from image: {e}")
            raise ValueError(f"Unable to extract text from image: {e}")

    @staticmethod
    def extract_text_from_txt(file_content: bytes) -> str:
        """Extract text from a plain text (.txt) file."""
        try:
            return file_content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return file_content.decode("latin-1")
            except Exception as e:
                logger.error(f"Failed to decode text file: {e}")
                raise ValueError(f"Unable to read text file: {e}")

    @classmethod
    def process_file(cls, file_content: bytes, filename: str, document_id: str = None) -> dict:
        """
        Process uploaded file and extract text and provenance based on file extension.
        Returns dict with extracted text, content_hash, document_id, and page provenance.
        """
        file_extension = os.path.splitext(filename)[1].lower()
        content_hash = hashlib.sha256(file_content).hexdigest()

        try:
            if file_extension in UNSUPPORTED_LEGACY_BINARY_EXTENSIONS:
                raise ValueError(
                    f"Legacy binary '{file_extension}' format is not supported. "
                    f"Please convert the file to OpenXML ('{file_extension}x') or PDF."
                )

            if file_extension == ".pdf":
                extractor = cls.extract_text_from_pdf
            elif file_extension == ".pptx":
                extractor = cls.extract_text_from_pptx
            elif file_extension == ".docx":
                extractor = cls.extract_text_from_docx
            elif file_extension in [".jpg", ".jpeg", ".png", ".bmp", ".tiff"]:
                extractor = cls.extract_text_from_image
            elif file_extension == ".txt":
                extractor = cls.extract_text_from_txt
            else:
                raise ValueError(f"Unsupported file type: {file_extension}")

            primary_error = None
            try:
                text = extractor(file_content)
            except Exception as e:
                primary_error = str(e)
                logger.info(f"Primary extraction failed for {filename}: {e}")
                text = ""

            from app.services.document_intelligence import DocumentIntelligence
            doc_intel = DocumentIntelligence.extract_rich_document(
                file_content, filename, document_id=document_id
            )
            if len((text or "").strip()) < 10:
                text = doc_intel.get("full_text", "")
            if not text or len(text.strip()) < 10:
                raise ValueError(
                    doc_intel.get("error")
                    or primary_error
                    or "No readable text could be extracted from the file."
                )

            return {
                "success": True,
                "text": text,
                "word_count": len(text.split()),
                "character_count": len(text),
                "file_type": file_extension,
                "filename": filename,
                "document_id": doc_intel.get("document_id"),
                "content_hash": content_hash,
                "headings": doc_intel.get("headings", []),
                "topics": doc_intel.get("topics", []),
                "definitions": doc_intel.get("definitions", []),
                "formulas": doc_intel.get("formulas", []),
                "elements": doc_intel.get("elements", []),
                "total_pages": doc_intel.get("total_pages", 1),
                "pages": doc_intel.get("pages", []),
            }

        except Exception as e:
            logger.error(f"Error processing file {filename}: {e}")
            return {
                "success": False,
                "error": str(e),
                "filename": filename,
                "content_hash": content_hash,
            }

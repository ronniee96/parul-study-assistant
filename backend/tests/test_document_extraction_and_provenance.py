import io
import unittest
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.services.document_processor import DocumentProcessor
from app.services.document_intelligence import DocumentIntelligence
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from PIL import Image
import docx
import pptx

class DocumentExtractionAndProvenanceTests(unittest.TestCase):
    def test_01_text_based_pdf_extraction_and_provenance(self):
        # Create a 2-page text PDF with ReportLab
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        
        # Page 1
        c.drawString(100, 700, "Parul University Examination Study Notes")
        c.drawString(100, 650, "Data Structures: A Stack is a LIFO linear data structure.")
        c.drawString(100, 600, "Operations on stack include Push, Pop, and Peek.")
        c.showPage()
        
        # Page 2
        c.drawString(100, 700, "Queue: A Queue is a FIFO linear data structure.")
        c.drawString(100, 650, "Operations on queue include Enqueue and Dequeue.")
        c.showPage()
        c.save()
        pdf_bytes = buf.getvalue()

        result = DocumentProcessor.process_file(pdf_bytes, "test_stack_queue.pdf")
        self.assertTrue(result["success"])
        self.assertIn("Stack is a LIFO", result["text"])
        self.assertIn("Queue is a FIFO", result["text"])
        self.assertEqual(result["total_pages"], 2)
        
        # Verify page numbers and provenance elements
        self.assertGreaterEqual(len(result["pages"]), 2)
        p1 = result["pages"][0]
        self.assertEqual(p1["page_number"], 1)
        self.assertIn("Stack", p1["text"])
        
        p2 = result["pages"][1]
        self.assertEqual(p2["page_number"], 2)
        self.assertIn("Queue", p2["text"])
        print("\n✓ Text-based PDF extraction and page provenance verified across 2 pages.")

    def test_02_scanned_unreadable_pdf_fails_cleanly(self):
        # Create a PDF with a pure blank/noise raster image (no text layer)
        img = Image.new('RGB', (200, 200), color=(255, 255, 255))
        img_buf = io.BytesIO()
        img.save(img_buf, format='JPEG')
        img_buf.seek(0)

        from reportlab.lib.utils import ImageReader
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        c.drawImage(ImageReader(img_buf), 100, 500, width=200, height=200)
        c.showPage()
        c.save()
        scanned_pdf_bytes = buf.getvalue()

        result = DocumentProcessor.process_file(scanned_pdf_bytes, "pure_scan.pdf")
        # Should NOT succeed with fake text; should return success: False with clear error
        self.assertFalse(result["success"])
        self.assertIn("error", result)
        print("✓ Scanned unreadable PDF produced clean error message (no hallucinated text).")

    def test_03_plain_text_extraction(self):
        sample = b"Operating Systems: Paging and Segmentation.\nPaging avoids external fragmentation."
        res = DocumentProcessor.process_file(sample, "os_notes.txt")
        self.assertTrue(res["success"])
        self.assertIn("Paging avoids external fragmentation", res["text"])
        print("✓ Plain text (.txt) extraction verified.")

    def test_04_docx_extraction(self):
        doc = docx.Document()
        doc.add_heading("Software Engineering", level=1)
        doc.add_paragraph("Agile methodology promotes adaptive planning and evolutionary development.")
        buf = io.BytesIO()
        doc.save(buf)
        res = DocumentProcessor.process_file(buf.getvalue(), "software_eng.docx")
        self.assertTrue(res["success"])
        self.assertIn("Agile methodology", res["text"])
        print("✓ Word document (.docx) extraction verified.")

    def test_05_pptx_extraction(self):
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Computer Architecture"
        slide.placeholders[1].text = "Pipelining hazards: Structural, Data, and Control hazards."
        buf = io.BytesIO()
        prs.save(buf)
        res = DocumentProcessor.process_file(buf.getvalue(), "architecture.pptx")
        self.assertTrue(res["success"])
        self.assertIn("Pipelining hazards", res["text"])
        print("✓ PowerPoint (.pptx) extraction verified.")

if __name__ == "__main__":
    unittest.main()
